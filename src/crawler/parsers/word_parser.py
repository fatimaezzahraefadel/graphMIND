"""Parseur Word (.docx) utilisant python-docx."""

from __future__ import annotations

import io

from src.shared.exceptions import CorruptedFileError
from src.shared.logging import get_logger
from src.crawler.parsers.base import BaseParser

logger = get_logger("word_parser")


class WordParser(BaseParser):
    """Extrait les paragraphes et tableaux d'un document Word .docx."""

    def can_parse(self, document_type: str) -> bool:
        return document_type == "word"

    def parse(self, content: bytes, source: str) -> str:
        """Extrait le texte d'un fichier .docx.

        Traite les paragraphes dans l'ordre puis les tableaux (cellules
        séparées par tabulation, lignes par saut de ligne).

        Args:
            content: Contenu binaire du .docx.
            source: URL ou chemin (pour les logs).

        Returns:
            Texte extrait.

        Raises:
            CorruptedFileError: Si le fichier ne peut pas être lu.
        """
        try:
            import docx
        except ImportError as exc:  # pragma: no cover
            raise CorruptedFileError(
                "python-docx non installé",
                component="word_parser",
            ) from exc

        try:
            doc = docx.Document(io.BytesIO(content))
        except Exception as exc:
            logger.warning("word_parse_error", source=source, error=str(exc))
            raise CorruptedFileError(
                f"Fichier Word corrompu ou illisible : {source}",
                component="word_parser",
                context={"source": source, "error": str(exc)},
            ) from exc

        parts: list[str] = []

        # Paragraphes
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                parts.append(text)

        # Tableaux : lignes avec cellules séparées par tabulation
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                # Dédupliquer les cellules fusionnées (python-docx les répète)
                seen: list[str] = []
                for cell in cells:
                    if not seen or cell != seen[-1]:
                        seen.append(cell)
                row_text = "\t".join(c for c in seen if c)
                if row_text:
                    parts.append(row_text)

        return "\n".join(parts)
